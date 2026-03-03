from __future__ import annotations

from pptx import Presentation

from core.document_model import DocumentObject


def parse(path: str) -> DocumentObject:
    prs = Presentation(path)
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                txt = shape.text.strip()
                if txt:
                    lines.append(txt)
    return DocumentObject(type="text", content=lines, metadata={"source": path})
