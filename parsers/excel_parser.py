from __future__ import annotations

from openpyxl import load_workbook

from core.document_model import DocumentObject


def parse(path: str) -> DocumentObject:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if c is None else str(c) for c in row])
    return DocumentObject(type="table", content=rows, metadata={"source": path, "sheet": ws.title})
